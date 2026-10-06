"""Generate synthetic DOCX evidence for the vendor-onboarding RAG demo.

All organizations, data, reports, and control statements in these files are fictional.
They are designed for software testing only and are not legal/compliance advice.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "sample-data" / "vendor-onboarding"
OUT.mkdir(parents=True, exist_ok=True)

BANK = "Acme Demo Bank"
VENDOR = "Northstar Demo Systems LLC"


def new_doc(title: str, code: str, version: str, owner: str, doc_type: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(42, 51, 45)
    for style_name, size, color in (("Title", 24, "235C45"), ("Heading 1", 16, "235C45"), ("Heading 2", 12, "315F47")):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FICTIONAL DEMO DOCUMENT · NOT A REAL BANK OR VENDOR RECORD")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(159, 91, 46)

    doc.add_heading(title, 0)
    meta = doc.add_table(rows=0, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Light Shading Accent 1"
    for label, value in (
        ("Document code", code),
        ("Version", version),
        ("Owner / author", owner),
        ("Document type", doc_type),
        ("Effective / submitted", "15 September 2026 (fictional demo date)"),
        ("Organization", BANK if owner.startswith(BANK) else VENDOR),
    ):
        cells = meta.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph(
        "Use only as synthetic portfolio-demo material. This document does not establish actual "
        "bank policy, regulatory compliance, or a real vendor certification."
    )
    return doc


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Light Shading Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(tbl.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = tbl.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value


def save(doc: Document, filename: str) -> None:
    doc.save(OUT / filename)


# 1. Bank's governing risk policy: the RAG should treat this as the primary decision source.
doc = new_doc(
    "Third-Party IT Risk Assessment Standard",
    "TPRM-STD-001",
    "1.0",
    f"{BANK} Technology Risk",
    "Internal risk standard",
)
doc.add_heading("1. Purpose and scope", level=1)
doc.add_paragraph(
    "This fictional standard describes how Acme Demo Bank assesses external technology providers "
    "before they receive bank data, connect to bank systems, or support an important bank service. "
    "It applies to cloud, software, managed-service, infrastructure, data-processing, and AI providers."
)
doc.add_heading("2. Inherent risk tiers", level=1)
table(doc, ["Tier", "Example trigger", "Review expectation"], [
    ["Low", "No bank data, no production access, and easy replacement.", "Business owner and Procurement review; confirm scope and contract owner."],
    ["Moderate", "Internal information or limited non-privileged access; service is replaceable.", "Business owner, Procurement, and Information Security review proportionate evidence."],
    ["High", "Customer or restricted data, privileged access, material subcontracting, or significant customer impact.", "Security, Privacy, Legal, Business Continuity, Procurement, and service owner review."],
    ["Critical", "Supports a core banking/payment function or an outage could materially disrupt essential services.", "High-tier reviews plus senior risk approval, tested exit/continuity plan, and enhanced monitoring."],
])
doc.add_paragraph(
    "A vendor is at least High when it processes customer identity data, government-ID images, "
    "authentication data, or other Restricted information, or when it can administer a production system. "
    "A high inherent tier cannot be reduced solely because a vendor has a security certificate."
)
doc.add_heading("3. Minimum evidence for an IT vendor", level=1)
bullets(doc, [
    "Service description, architecture/data-flow diagram or equivalent, data categories, hosting regions, system access, and subcontractor list.",
    "Completed security and privacy due-diligence questionnaire, including exceptions and evidence references.",
    "Independent security assurance appropriate to the tier, such as a scoped SOC 2 Type II report or ISO 27001 certificate and Statement of Applicability. Review report scope, period, exceptions, and complementary user-entity controls.",
    "Recent independent penetration-test executive summary and remediation status. Do not request unrestricted raw test details unless an approved reviewer needs them.",
    "Business-continuity/disaster-recovery summary, latest exercise results, recovery targets, and security-incident response/notification process.",
    "For personal or Restricted data: privacy terms, retention/deletion schedule, subprocessors, data locations, and permitted-use statement.",
])
doc.add_heading("4. Review and decision rules", level=1)
numbered(doc, [
    "The business owner documents the service purpose, criticality, data, system access, and intended launch date.",
    "Technology Risk assigns an inherent tier before considering controls.",
    "Control owners assess the evidence, record gaps with source references, and distinguish verified evidence from vendor statements.",
    "A High or Critical vendor cannot go live with an unresolved material gap unless an authorized, time-limited risk exception is approved.",
    "The final decision, conditions, owner, expiry date, and monitoring plan are recorded in the vendor case.",
])
doc.add_heading("5. Evidence quality", level=1)
doc.add_paragraph(
    "A vendor questionnaire is a self-attestation, not independent assurance. An expired certificate, "
    "out-of-scope report, missing report, or unremediated finding must be identified as a gap rather than "
    "described as evidence of compliance. The reviewer must not infer that an unprovided control exists."
)
save(doc, "01_TPRM_STD_001_Third_Party_IT_Risk_Assessment_Standard.docx")

# 2. Bank's control baseline used to test vendor responses.
doc = new_doc(
    "Third-Party Information Security and Data Handling Requirements",
    "SEC-STD-004",
    "1.0",
    f"{BANK} Information Security",
    "Internal security standard",
)
doc.add_heading("1. Data classification", level=1)
table(doc, ["Class", "Demo definition", "Vendor handling"], [
    ["Public", "Approved for public release.", "No special bank restrictions beyond contract and integrity requirements."],
    ["Internal", "Non-public operational information with limited impact if exposed.", "Use approved access, secure transmission, and documented deletion."],
    ["Confidential", "Sensitive business, employee, or customer information.", "Need-to-know access, encryption, logging, and approved subprocessors."],
    ["Restricted", "Customer identity documents, account-linked data, credentials, or highly sensitive information.", "High-risk review, strict least privilege, strong encryption, MFA, monitoring, and approved data-use/retention terms."],
])
doc.add_heading("2. Baseline controls for vendors", level=1)
table(doc, ["Control area", "Acme Demo Bank requirement"], [
    ["Encryption", "Protect bank data in transit using TLS 1.2 or stronger and at rest using an industry-accepted strong encryption method."],
    ["Identity and access", "Require MFA for all workforce access to production or bank data; use unique accounts, least privilege, and time-limited privileged access. Shared administrator accounts are not permitted."],
    ["Logging", "Keep security-relevant access and administrative logs for at least 365 days for High and Critical services, with protected access and review."],
    ["Vulnerability response", "Fix Critical vulnerabilities within 72 hours or document compensating safeguards and an approved exception; fix High vulnerabilities within 15 calendar days."],
    ["Security testing", "Complete an independent penetration test at least annually and track findings to closure. Material findings must be disclosed to the bank reviewer."],
    ["Incident notice", "Notify the bank within 24 hours after becoming aware of a suspected or confirmed incident affecting bank data or the contracted service. Initial notice may be incomplete and updated as facts are confirmed."],
    ["Data use and AI", "Do not use bank data, prompts, outputs, or files to train or improve a shared/general model. Any different use requires written bank approval before processing."],
    ["Subprocessors", "Maintain a current subprocessor list, apply equivalent controls, and notify the bank at least 30 days before a material change, subject to contract approval rights."],
    ["Deletion and return", "Delete active bank data within 30 days after contract end or bank request; backups may age out within 90 days if isolated, protected, and not restored except for recovery."],
])
doc.add_heading("3. Evidence and exceptions", level=1)
doc.add_paragraph(
    "Vendor claims must be tied to evidence, scope, and date. A gap is not closed by a general policy statement. "
    "A compensating control must reduce the specific risk, have a named owner, and be verified. High-risk "
    "exceptions require written approval and a maximum initial duration of 90 days."
)
doc.add_paragraph(
    "This fictional baseline is intentionally a demo control set; it is not a statement of any real bank's "
    "requirements or a substitute for applicable law, regulation, or approved institutional policy."
)
save(doc, "02_SEC_STD_004_Third_Party_Security_and_Data_Handling_Requirements.docx")

# 3. Bank's resilience and exception rules, deliberately compared to vendor claims below.
doc = new_doc(
    "Third-Party Resilience, Recovery, and Incident Requirements",
    "BCM-STD-006",
    "1.0",
    f"{BANK} Business Continuity and Technology Risk",
    "Internal resilience standard",
)
doc.add_heading("1. Recovery objectives", level=1)
table(doc, ["Service tier", "Maximum recovery time (RTO)", "Maximum data loss window (RPO)", "Exercise"], [
    ["Moderate", "24 hours", "8 hours", "Documented recovery test at least every 24 months."],
    ["High", "4 hours", "1 hour", "End-to-end recovery exercise at least annually; retain results and remediation."],
    ["Critical", "1 hour", "15 minutes", "Annual end-to-end exercise plus a documented failover/exit test."],
])
doc.add_paragraph(
    "RTO is the target time to restore the service. RPO is the maximum acceptable period of data that may need to be recreated after recovery. Vendor objectives must support the bank's service-level needs, not merely the vendor's standard tier."
)
doc.add_heading("2. Incident response", level=1)
bullets(doc, [
    "The vendor must maintain a named 24/7 security contact for High and Critical services.",
    "Initial notice is due within 24 hours after awareness of an incident affecting bank data or the contracted service; the vendor must provide updates as facts are confirmed.",
    "The notice must identify known impact, affected service/data, containment steps, contact, and next update time. Root-cause details may follow when established.",
    "The vendor must preserve relevant evidence, cooperate with the bank, and not make public statements naming the bank without authorization except where law requires.",
])
doc.add_heading("3. Continuity evidence", level=1)
bullets(doc, [
    "Provide the most recent recovery exercise date, scope, achieved RTO/RPO, failures, and corrective-action owners/dates.",
    "Identify single points of failure, key hosting/subprocessor dependencies, backup protection, and manual workarounds.",
    "A tabletop exercise alone does not demonstrate technical recovery of a production service; describe what was actually restored or failed over.",
])
doc.add_heading("4. Exceptions", level=1)
doc.add_paragraph(
    "A vendor exception to a recovery target or notice period requires service-owner justification, Business Continuity and Technology Risk review, compensating measures, an approved contract position, a named owner, and an expiry date. Critical services require senior risk approval."
)
save(doc, "03_BCM_STD_006_Third_Party_Resilience_and_Incident_Requirements.docx")

# 4. Synthetic vendor intake and data-flow responses.
doc = new_doc(
    "IT Vendor Service and Data-Flow Questionnaire — Northstar Demo Systems",
    "VND-NSTAR-INTAKE-001",
    "1.0",
    f"{VENDOR} (fictional respondent)",
    "Vendor due-diligence questionnaire",
)
doc.add_heading("1. Service overview", level=1)
doc.add_paragraph(
    "Service: Northstar Verify, a hosted identity-verification service proposed for Acme Demo Bank's online account-opening workflow. It accepts an applicant's identity details and document images and returns a verification result and case reference."
)
doc.add_paragraph("Business owner: Digital Account Opening. Planned use: production. Availability is material to new-account onboarding; the service is not the bank's core deposit ledger.")
doc.add_heading("2. Data and access", level=1)
table(doc, ["Question", "Fictional vendor response"], [
    ["Data processed", "Applicant name, date of birth, residential address, email, government-ID image, selfie image, verification result, and bank-generated case ID."],
    ["Classification requested", "Northstar classifies images and identity fields as confidential customer data. Acme Demo Bank treats them as Restricted."],
    ["Storage region", "Primary production data is hosted in a US region. Disaster-recovery copies are in a second US region. No EU or Asia storage is declared."],
    ["Bank access", "The service uses an API integration and a service account. Northstar support staff may request time-limited remote access for troubleshooting. The vendor reports that support access is approved per case."],
    ["Subprocessors", "Nimbus Harbor Cloud LLC provides hosting; Lantern OCR Services LLC provides document-text extraction. These names and relationships are fictional."],
    ["AI/model use", "Lantern OCR may retain request content for up to 30 days for abuse investigation. Northstar states that customer content is not used to train a shared model, but the vendor has not supplied a contract clause confirming this restriction."],
    ["Deletion", "Northstar states that active data is deleted within 45 days after termination. Encrypted backups age out within 90 days."],
])
doc.add_heading("3. Service and security statements", level=1)
bullets(doc, [
    "TLS is used for external API connections; the vendor states that stored production data is encrypted at rest.",
    "Workforce access uses SSO and MFA for most systems. Emergency local administrator accounts are described in the separate security response.",
    "The vendor has not provided a current independent SOC 2 Type II report or ISO 27001 certificate with this questionnaire.",
    "The vendor asks Acme Demo Bank to review its security summary and test evidence before production approval.",
])
doc.add_heading("4. Questionnaire limitations", level=1)
doc.add_paragraph(
    "This is a vendor self-attestation created for a fictional demo. Statements are not independently verified. Missing documents must be treated as missing evidence, not proof that a control is absent or present."
)
save(doc, "04_VND_NSTAR_INTAKE_001_Service_and_Data_Flow_Questionnaire.docx")

# 5. Synthetic security evidence summary; explicitly not a fabricated SOC 2 report.
doc = new_doc(
    "Vendor Security Assurance and Testing Summary — Northstar Demo Systems",
    "VND-NSTAR-SEC-002",
    "1.0",
    f"{VENDOR} (fictional respondent)",
    "Vendor self-attestation and evidence index",
)
doc.add_heading("Important status", level=1)
doc.add_paragraph(
    "This file is a fictional vendor-prepared summary for a demo. It is NOT a SOC 2 report, ISO 27001 certificate, penetration-test report, or independent auditor opinion. No real assurance report is included."
)
doc.add_heading("1. Assurance documents", level=1)
table(doc, ["Evidence item", "Status stated by vendor", "Reviewer note"], [
    ["SOC 2 Type II", "Audit fieldwork is in progress; report not yet issued or provided.", "Independent assurance is missing for this review. Do not mark SOC 2 as passed."],
    ["ISO 27001", "No certificate provided.", "No certification claim is made in this packet."],
    ["Penetration test", "External test dated 30 June 2025; executive summary only; one Medium finding remains open with target closure 31 October 2026.", "The assessment is over 12 months old at the fictional review date. Obtain an updated test and closure evidence."],
])
doc.add_heading("2. Selected control responses", level=1)
table(doc, ["Control", "Fictional vendor response", "Evidence status"], [
    ["Privileged MFA", "SSO and MFA are enforced for workforce users. Two emergency local administrator accounts are excluded from MFA; access alerts are reviewed after use.", "Exception described, but no compensating-control test supplied."],
    ["Access review", "Managers review production access twice per year.", "No completed review sample attached."],
    ["Security logs", "Administrative and authentication logs are retained for 180 days.", "Vendor statement only; bank standard requires 365 days for High service."],
    ["Critical vulnerability remediation", "Target is seven calendar days after validation; emergency fixes may be deployed sooner.", "Bank standard requires 72 hours or an approved exception."],
    ["Encryption", "Vendor states TLS 1.2 or stronger in transit and encryption at rest for production storage.", "Configuration evidence not attached."],
    ["AI training use", "Vendor says customer content is not used to train a shared model.", "Contractual confirmation and Lantern OCR retention terms not attached."],
])
doc.add_heading("3. Open items acknowledged by vendor", level=1)
bullets(doc, [
    "Provide the independent SOC 2 Type II report when issued, including scope, period, exceptions, and complementary user-entity controls.",
    "Provide a current independent penetration-test executive summary and remediation evidence for the open Medium finding.",
    "Provide evidence for emergency-account MFA compensating controls and the latest production-access review.",
    "Confirm the 24-hour incident-notice term, 365-day log availability, 72-hour critical patch path, AI data-use restriction, and 30-day active-data deletion in contract or approved exception documents.",
])
doc.add_heading("4. Reviewer caution", level=1)
doc.add_paragraph(
    "This summary is self-reported and synthetic. It must not be described as an audit report or used to claim real-world certification."
)
save(doc, "05_VND_NSTAR_SEC_002_Security_Assurance_and_Test_Summary.docx")

# 6. Synthetic BCP/DR and incident response evidence with deliberate measurable gaps.
doc = new_doc(
    "Business Continuity and Incident Response Summary — Northstar Demo Systems",
    "VND-NSTAR-BCM-003",
    "1.0",
    f"{VENDOR} (fictional respondent)",
    "Vendor resilience questionnaire response",
)
doc.add_heading("1. Recovery objectives", level=1)
table(doc, ["Measure", "Vendor-stated target", "Evidence supplied"], [
    ["Recovery time objective (RTO)", "8 hours for Northstar Verify.", "Target stated; no service-specific bank workload mapping supplied."],
    ["Recovery point objective (RPO)", "4 hours for applicant case data.", "Target stated; restoration sample not supplied."],
    ["Last recovery exercise", "12 August 2025; tabletop and partial failover.", "Summary only. Full end-to-end restoration was not demonstrated."],
    ["Exercise result", "A DNS cutover issue delayed restoration; corrective work is marked in progress.", "No closure evidence or retest date supplied."],
])
doc.add_heading("2. Dependencies and continuity", level=1)
bullets(doc, [
    "The service depends on Nimbus Harbor Cloud LLC for hosting and Lantern OCR Services LLC for document extraction.",
    "The vendor has not supplied a recovery test for loss of the OCR subprocessor or the primary cloud region.",
    "A manual review queue is available for limited traffic, but staffing capacity and maximum duration are not specified.",
    "Backups are encrypted according to the vendor statement; backup restoration evidence was not attached.",
])
doc.add_heading("3. Security-incident response", level=1)
table(doc, ["Topic", "Fictional vendor response"], [
    ["Customer notification", "Contract template permits notice within 72 hours after confirming an incident affecting customer data."],
    ["24/7 contact", "A security mailbox is monitored continuously; a named escalation roster was not provided."],
    ["Initial notice content", "The vendor expects to provide known affected services and containment status; the initial template was not attached."],
    ["Customer cooperation", "The vendor states it will investigate and provide updates; contractual evidence and forensic-support scope were not supplied."],
])
doc.add_heading("4. Required follow-up for Acme Demo Bank", level=1)
bullets(doc, [
    "High-tier bank target is RTO 4 hours and RPO 1 hour; vendor-stated targets are 8 hours and 4 hours.",
    "High-tier recovery testing must be annual and demonstrate technical restoration; the supplied exercise is over 12 months old and only partially demonstrates failover.",
    "The bank standard requires initial incident notice within 24 hours; the vendor template allows up to 72 hours after confirmation.",
    "Obtain a retest plan and closure evidence for the DNS cutover issue, subprocessor outage coverage, backup restoration, and incident escalation contacts.",
])
doc.add_paragraph(
    "All responses and organizations in this file are synthetic. It is designed to test evidence comparison and gap identification, not to represent a real vendor."
)
save(doc, "06_VND_NSTAR_BCM_003_Business_Continuity_and_Incident_Response_Summary.docx")

print(f"Created 6 fictional vendor-onboarding DOCX files in {OUT}")
