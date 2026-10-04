"""Create fictional Acme procurement policy DOCX for local development."""
from pathlib import Path

from docx import Document


OUTPUT_PATH = (
    Path(__file__).resolve().parents[1]
    / "sample-data"
    / "source-documents"
    / "Acme_Procurement_Policy.docx"
)


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Acme Technologies LLC — Procurement Policy", level=0)
    doc.add_paragraph(
        "NOTICE: Fictional document for software demonstration only; not legal advice."
    )

    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "This document outlines standard procurement procedures for internal testing and system demonstration."
    )

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph(
        "Applies to all simulated vendor onboarding and invoice extraction workflows."
    )

    doc.add_heading("3. Vendor Review", level=1)
    doc.add_paragraph(
        "[FICTIONAL RULE] All vendors must provide verifiable tax identification numbers and clear sanctions screening before onboarding."
    )

    doc.add_heading("4. Approval", level=1)
    doc.add_paragraph(
        "[FICTIONAL RULE] Invoices exceeding $10,000 automatically trigger a Human-in-the-Loop review queue."
    )

    doc.add_heading("5. Recordkeeping", level=1)
    doc.add_paragraph(
        "All audit logs and transaction states must be archived in PostgreSQL for a minimum of 7 years."
    )

    doc.save(OUTPUT_PATH)
    print(f"Created fictional sample document: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
